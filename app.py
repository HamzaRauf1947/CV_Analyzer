import io
import re
from collections import Counter

import streamlit as st
import pdfplumber
from docx import Document
import spacy
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
import language_tool_python

st.set_page_config(page_title="AI CV Analyzer", page_icon="📄", layout="wide")

st.markdown("""
<style>
.main {background-color: #f7f9fc;}
.block-container {padding-top: 2rem; max-width: 1200px;}
.hero {padding: 1.5rem; border-radius: 18px; background: linear-gradient(120deg,#14213d,#315caa); color:white;}
.small-muted {color:#64748b;}
</style>
""", unsafe_allow_html=True)

@st.cache_resource
def load_nlp():
    try:
        return spacy.load("en_core_web_sm")
    except OSError:
        return None

@st.cache_resource
def load_grammar_tool():
    try:
        return language_tool_python.LanguageTool("en-US")
    except Exception:
        return None

def extract_text(uploaded_file):
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()
    if name.endswith(".pdf"):
        pieces = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                pieces.append(page.extract_text() or "")
        return "\n".join(pieces).strip()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts).strip()
    raise ValueError("Please upload a PDF or DOCX file.")

def section_presence(text):
    lower = text.lower()
    patterns = {
        "Contact information": r"\b(email|phone|contact|linkedin|github)\b",
        "Professional summary": r"\b(summary|professional profile|objective)\b",
        "Skills": r"\b(skills|technical skills|core competencies)\b",
        "Experience": r"\b(experience|employment|work history|internship)\b",
        "Education": r"\b(education|academic background|qualification)\b",
        "Projects": r"\b(projects|personal projects|academic projects)\b",
    }
    return {label: bool(re.search(pattern, lower)) for label, pattern in patterns.items()}

def keyword_match(resume, job_description):
    if not job_description.strip():
        return None, [], []
    vectorizer = TfidfVectorizer(stop_words="english", ngram_range=(1, 2))
    try:
        matrix = vectorizer.fit_transform([resume, job_description])
        similarity = float(cosine_similarity(matrix[0:1], matrix[1:2])[0][0] * 100)
    except ValueError:
        return 0.0, [], []
    resume_words = set(re.findall(r"[a-zA-Z][a-zA-Z+#.-]{1,}", resume.lower()))
    job_words = set(re.findall(r"[a-zA-Z][a-zA-Z+#.-]{1,}", job_description.lower()))
    stop = {"with", "from", "that", "this", "will", "your", "have", "are", "for",
            "and", "the", "you", "our", "their", "they", "who", "what", "when",
            "where", "can", "all", "use", "using", "ability", "work", "role",
            "team", "good", "strong", "years", "year", "required", "preferred"}
    candidates = [w for w in job_words if w not in stop and len(w) > 2]
    missing = sorted(w for w in candidates if w not in resume_words)
    matched = sorted(w for w in candidates if w in resume_words)
    return similarity, missing[:30], matched[:30]

def grammar_check(text, tool):
    if not tool:
        return [], "LanguageTool could not be started. Grammar checking is unavailable; other analysis can still run."
    try:
        matches = tool.check(text[:20000])
        results = []
        for m in matches[:50]:
            context = text[max(0, m.offset - 45):min(len(text), m.offset + m.error_length + 45)].replace("\n", " ")
            results.append({
                "Issue": m.message,
                "Context": context,
                "Suggested correction": ", ".join(m.replacements[:3]) if m.replacements else "Review manually",
                "Rule": m.rule_id
            })
        return results, None
    except Exception as exc:
        return [], f"Grammar checker error: {exc}"

def weak_phrases(text):
    phrases = [
        ("was responsible for", "Use a specific action verb such as managed, developed, or coordinated—if accurate."),
        ("responsible for", "Describe what you actually did and the result."),
        ("helped with", "Specify your contribution and outcome."),
        ("worked on", "Use a precise verb such as designed, tested, analyzed, or implemented—if accurate."),
        ("hard worker", "Replace with evidence of achievements or skills."),
        ("team player", "Give a brief example of collaboration and its result."),
        ("duties included", "Describe accomplishments rather than only listing duties."),
    ]
    low = text.lower()
    found = []
    for phrase, tip in phrases:
        if phrase in low:
            found.append((phrase, tip))
    return found

def calculate_score(sections, grammar_count, similarity):
    section_score = sum(sections.values()) / max(1, len(sections)) * 40
    grammar_score = max(0, 30 - min(grammar_count, 30))
    keyword_score = similarity * 0.30 if similarity is not None else 15
    return round(max(0, min(100, section_score + grammar_score + keyword_score)))

st.markdown("""
<div class="hero">
  <h1>📄 AI-Powered CV / Resume Analyzer</h1>
  <p>Natural Language Processing for resume quality checks and job-description alignment</p>
</div>
""", unsafe_allow_html=True)
st.write("Upload a CV to check its sections, writing issues, readability indicators, and keyword alignment.")
st.caption("Privacy note: this prototype analyzes files in the running app session and does not intentionally save uploaded CVs. Avoid uploading sensitive personal information to public deployments.")

with st.sidebar:
    st.header("Analysis options")
    st.write("Supported formats: PDF and DOCX")
    st.write("NLP: spaCy (when installed)")
    st.write("Grammar: LanguageTool")
    st.write("Matching: TF-IDF + cosine similarity")
    st.divider()
    st.caption("This tool provides guidance, not a guarantee of ATS performance or hiring outcomes.")

uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx"])
job_description = st.text_area("Paste the target job description (optional)", height=180,
                               placeholder="Paste the job advertisement here to compare relevant terms and skills.")

if uploaded:
    try:
        resume_text = extract_text(uploaded)
    except Exception as exc:
        st.error(f"Could not read the file: {exc}")
        st.stop()

    if not resume_text:
        st.error("No readable text was extracted. If this is a scanned PDF, run OCR on it first or upload a text-based PDF/DOCX.")
        st.stop()

    st.success(f"Text extracted successfully: {len(resume_text.split())} words.")
    with st.expander("Preview extracted text"):
        st.text_area("Extracted resume text", resume_text, height=240)

    if st.button("Analyze Resume", type="primary", use_container_width=True):
        with st.spinner("Analyzing your resume..."):
            sections = section_presence(resume_text)
            nlp = load_nlp()
            grammar_tool = load_grammar_tool()
            grammar_results, grammar_warning = grammar_check(resume_text, grammar_tool)
            similarity, missing, matched = keyword_match(resume_text, job_description)
            weak = weak_phrases(resume_text)
            score = calculate_score(sections, len(grammar_results), similarity)

            st.session_state["analysis"] = {
                "sections": sections, "grammar": grammar_results, "grammar_warning": grammar_warning,
                "similarity": similarity, "missing": missing, "matched": matched,
                "weak": weak, "score": score, "text": resume_text,
                "nlp_available": nlp is not None
            }

if "analysis" in st.session_state:
    result = st.session_state["analysis"]
    st.divider()
    st.header("Analysis report")
    c1, c2, c3 = st.columns(3)
    c1.metric("Resume guidance score", f'{result["score"]}/100')
    c2.metric("Grammar issues found", len(result["grammar"]))
    c3.metric("Words extracted", len(result["text"].split()))
    st.caption("The score is a simple project-defined heuristic, not a validated hiring or ATS score.")

    left, right = st.columns(2)
    with left:
        st.subheader("Section check")
        for section, present in result["sections"].items():
            st.write(("✅ " if present else "⚠️ ") + section + (": detected" if present else ": not detected"))
        st.info("Section detection searches for common headings and may miss creative or unusual heading names.")
    with right:
        st.subheader("Job-description alignment")
        if result["similarity"] is None:
            st.write("Add a job description to calculate text similarity and keyword differences.")
        else:
            st.metric("TF-IDF cosine similarity", f'{result["similarity"]:.1f}%')
            st.write("**Terms found in both texts (approximate):**")
            st.write(", ".join(result["matched"]) or "No simple keyword overlap detected.")
            st.write("**Potentially missing terms:**")
            st.write(", ".join(result["missing"]) or "No missing terms detected by this basic comparison.")
            st.caption("Keyword overlap is approximate and can include generic words. Review each term for relevance before editing your CV.")

    st.subheader("Grammar and spelling suggestions")
    if result["grammar_warning"]:
        st.warning(result["grammar_warning"])
    if result["grammar"]:
        st.dataframe(result["grammar"], use_container_width=True, hide_index=True)
    elif not result["grammar_warning"]:
        st.success("No issues were returned by the grammar checker for the analyzed text.")

    st.subheader("Phrasing suggestions")
    if result["weak"]:
        for phrase, tip in result["weak"]:
            st.write(f"- **“{phrase}”**: {tip}")
    else:
        st.write("No phrases from the built-in weak-phrase list were detected.")

    st.subheader("Readability indicators")
    text = result["text"]
    sentences = [s for s in re.split(r"[.!?]+", text) if s.strip()]
    avg_words = len(text.split()) / max(1, len(sentences))
    st.write(f"- Approximate average sentence length: **{avg_words:.1f} words**.")
    if avg_words > 25:
        st.write("- Some sentences may be long; consider splitting them where it improves clarity.")
    else:
        st.write("- Average sentence length is not especially long by this simple measure.")
    st.caption("This prototype does not yet reliably detect passive voice. Sentence length is only a rough readability indicator.")

    if result["nlp_available"]:
        st.success("spaCy English model is available. NLP model loaded successfully.")
    else:
        st.warning("spaCy English model is not installed. The app's rule-based checks and TF-IDF matching can still run.")

    report_lines = [
        "AI-POWERED CV / RESUME ANALYZER — REPORT",
        "=" * 45,
        f"Guidance score: {result['score']}/100",
        f"Words extracted: {len(result['text'].split())}",
        f"Grammar issues returned: {len(result['grammar'])}",
        "",
        "SECTION CHECK:"
    ]
    report_lines += [f"- {k}: {'Detected' if v else 'Not detected'}" for k, v in result["sections"].items()]
    report_lines += ["", "PHRASING SUGGESTIONS:"]
    report_lines += [f"- {p}: {tip}" for p, tip in result["weak"]] or ["- None detected by built-in list."]
    if result["similarity"] is not None:
        report_lines += ["", f"TF-IDF cosine similarity: {result['similarity']:.1f}%",
                         "Potentially missing terms: " + (", ".join(result["missing"]) or "None detected")]
    report_lines += ["", "Note: Automated guidance only; review suggestions manually."]
    st.download_button("Download summary report (.txt)", "\n".join(report_lines),
                       file_name="cv_analysis_report.txt", mime="text/plain")
else:
    st.info("Upload a CV, optionally paste a job description, then click **Analyze Resume**.")
